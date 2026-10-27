# -*- coding: utf-8 -*-
"""生成重构后的创新奖申报材料（通用方案本体 + 编码通(新)案例）"""
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


FONT_BODY = '仿宋_GB2312'
FONT_HEAD = '黑体'
FONT_MONO = 'Consolas'


def set_run_font(run, name, size_pt, bold=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(text)
    set_run_font(r, FONT_HEAD, 22, bold=True)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_font(r, FONT_HEAD, 16, bold=True)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, FONT_HEAD, 14, bold=True)
    return p


def add_body(doc, runs):
    """runs: list of (text, bold) tuples"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for text, bold in runs:
        r = p.add_run(text)
        set_run_font(r, FONT_BODY, 12, bold=bold)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, FONT_BODY, 10.5, bold=False)
    return p


def add_image(doc, path, width_inches, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_inches))
    if caption:
        add_note(doc, caption)


def add_placeholder(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_font(r, FONT_BODY, 11, bold=True)
    return p


def add_mono_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, FONT_MONO, 10, bold=False)
    return p


def add_bullet(doc, runs):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(8)
    for text, bold in runs:
        r = p.add_run(text)
        set_run_font(r, FONT_BODY, 12, bold=bold)
    return p


def add_table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = 'Table Grid'
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, FONT_BODY, 11, bold=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, FONT_BODY, 11, bold=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def build():
    doc = docx.Document()

    add_title(doc, '智能技术支持问答体系——面向低代码产品矩阵的可复制方案\n(核心落地案例：编码通(新))')

    # ============ 一、定义问题 ============
    add_h1(doc, '一、定义问题')

    add_h2(doc, '1.1 行业趋势与通用诉求')
    add_body(doc, [('随着大语言模型技术的成熟，结合企业私有知识的智能问答正在成为技术支持的主流范式。无论是面向内部员工还是面向外部客户的技术支持场景，"人工客服+文档检索"的传统模式普遍面临规模化瓶颈——重复性问题的解答边际成本居高不下，而真正需要经验判断的复杂问题又容易被淹没其中，难以获得足够关注。', False)])

    add_h2(doc, '1.2 技术支持场景的共性特征')
    add_body(doc, [('通过对实际技术支持咨询记录的系统性梳理，项目组提炼出三条具有普遍适用性的共性特征，这些特征并不局限于某一具体平台或场景：', False)])
    add_body(doc, [('特征一：用户问题天然分层。', True), ('用户咨询并非同质化的"一个问题一个答案"，而是天然可以划分为四种性质截然不同的类型：', False)])

    add_table(doc,
              ['问题类型', '特征描述', '占比'],
              [
                  ['确定性问题', '规则明确、答案固定，可通过标准知识直接解答', '【待补充】'],
                  ['推理性问题', '解决思路/方法固定，但需结合用户具体业务场景、数据结构生成针对性答案', '【待补充】'],
                  ['场景可行性评估', '需协助用户判断某一业务场景能否实现，涉及复杂业务判断，超出规则化、模板化应答的能力边界', '【待补充】'],
                  ['需求/BUG反馈', '用户对产品功能提出的改进诉求或缺陷反馈，非"问答"性质', '【待补充】'],
              ])
    add_note(doc, '说明：问题分类标准由项目组制定后，通过AI辅助分类的方式对历史咨询记录进行统计，具体占比数据将在规则确定后补充；该分类标准与占比统计方法与具体业务逻辑无关，可直接套用于其他技术支持场景。')

    add_body(doc, [('特征二：响应速度与答案质量的权衡。', True), ('追求极致的响应速度（如简单的关键词匹配）往往意味着答案质量的妥协；而追求答案的精准与个性化（如深度推理）又天然伴随着更高的计算成本和响应延迟，如何在两者间取得平衡，是任何AI问答系统都绕不开的设计难题。', False)])
    add_body(doc, [('特征三：问答闭环离不开人工参与。', True), ('机器既无法独立应答所有问题——复杂场景判断等超出当前AI能力边界的问题，最终仍需人工兜底；也无法独立完成自身能力的迭代——知识库和技能库的内容补充，同样依赖人工的判断与操作。这意味着任何技术支持问答体系要真正规模化运转，都必须为人工预留合理的参与环节，而非追求"完全无人化"。', False)])

    # ============ 二、方案本体 ============
    add_h1(doc, '二、方案本体：一套可复制的智能技术支持问答体系')
    add_body(doc, [('本章节介绍的是一套面向上述三条共性特征设计的', False), ('通用技术支持问答体系', True), ('，核心引擎、Agent协同模式等通用能力与知识库内容、SKILL清单等具体业务数据层彻底解耦，理论上可迁移到任意技术支持问答场景（实际落地与复用情况详见三、实践成效）。', False)])

    add_h2(doc, '2.1 设计理念：两条核心价值主张')
    add_body(doc, [('本方案的设计理念由两条核心价值主张支撑，分别回应1.2中的特征一与特征二；特征三（问答闭环离不开人工参与）不是独立的设计主张，而是体现在2.2-2.5各环节对人工角色边界的具体安排中：', False)])
    add_body(doc, [('价值主张一：问题分层应答，而非单一维度的问答功能。', True), ('基于"用户问题天然分层"的洞察，设计"检索—推理—人工"能力分级响应、且具备知识自我闭环能力的体系，实现从"人工被动答疑"到"AI主动分层拦截、人工高价值兜底"的转变。', False)])
    add_body(doc, [('价值主张二：计算资源利用效率与用户体验的动态平衡，而非非此即彼的取舍。', True), ('通过上游单点精确决策、按需触发下游执行、状态组合驱动的终态判断与分阶段渐进交付等机制，把有限的计算资源集中投入到真正有价值的推理环节上。', False)])

    add_h2(doc, '2.2 通用架构概览')
    add_body(doc, [('系统按"预处理→核心应答→渐进交付与人工协同→知识自增长"四个阶段构建通用链路，各阶段的编排逻辑与具体业务平台解耦，接入新场景只需替换知识库、SKILL清单、文档库等数据层。架构中的人工环节，是当前AI能力尚无法完全自主覆盖复杂判断与知识缺口填补的务实安排，而非设计追求的理念本身。', False)])
    add_image(doc, 'architecture_diagram.png', 5.0, caption='图1 整体架构图：两个人工入口、核心应答链路与知识自增长闭环')

    add_h2(doc, '2.3 问题预处理：查询标准化与提问有效性过滤')
    add_body(doc, [('用户提问口语化、表达随意，直接检索容易"文不对题"。系统先通过大模型将原始问题改写为书面化标准问句，再进入检索加工。同时引入提问有效性过滤：规则层过滤明显无效内容，语义层判断是否包含真实的问题、报错或阻碍，只有通过双层过滤的提问才计入熔断计时（具体规则不做展开，避免被针对性绕过）。', False)])

    add_h2(doc, '2.4 核心引擎：三级递进应答与人工入口动态管控')
    add_body(doc, [('（1）三级应答能力', True)])
    add_body(doc, [('第一级·知识库检索：', True), ('对接企业知识库做向量检索，知识库内容来源可插拔，与具体业务无关。检索完成后，系统才基于当前问题发起技能判断（见(2)），两者严格先后、不并行——终态判断依赖知识库答案这一前置结果。', False)])
    add_body(doc, [('第二级·Agent技能推理：', True), ('针对需结合用户具体业务场景生成定制化答案的问题，由双Agent解耦架构（见(2)）判断并执行可用技能。', False)])
    add_body(doc, [('第三级·文档推荐：', True), ('前两级均未命中时，基于关键词匹配推荐相关文档，避免"无法回答"的空结果。', False)])
    add_body(doc, [('机器人答案同步展示知识来源支持溯源，并支持同一对话内的多轮追问。', False)])

    add_body(doc, [('（2）双Agent解耦架构：上游单点决策，下游纯执行', True)])
    add_body(doc, [('技能推理为何要拆成两个Agent，而非一个Agent"判断+执行"一气呵成？根本原因在于(3)所述的终态判断与渐进交付机制，需要尽早拿到"是否存在可用技能"这一信号，才能决定是提示用户"更精准答案在路上"，还是直接把知识库答案当终态呈现——缺了这个信号，要么误把过渡态答案当终态，要么让用户盲目等待产生焦虑。但"判断有没有技能"是轻量语义匹配，"调用技能生成答案"却涉及实际执行、耗时高得多：若两者耦合在一个Agent里顺序完成，就必须等执行完才能拿到判断信号，渐进交付也就无从谈起。因此拆分为上游（快速产出信号）与下游（专注执行）两个Agent，这是渐进交付机制得以成立的工程前提。', False)])
    add_bullet(doc, [('上游·技能决策Agent：', True), ('唯一承担语义判断职责——识别单一或最多2个子需求，每个子需求基于"技能描述所覆盖业务范畴是否与需求相符"精确匹配1个最终技能，产出下游可直接执行、无需二次判断的技能清单。', False)])
    add_bullet(doc, [('下游·技能执行Agent：', True), ('不做二次语义判断，按清单逐一执行（容忍技能名称≥90%相似度的表述差异），仅在技能确实不存在或调用无有效结果时判定无法执行。', False)])
    add_body(doc, [('该职责划分由真实运行反馈驱动：早期"上下游各自判断"的两段式设计中，下游会对上游已确认的技能重新产生怀疑、倾向保守拒绝执行，多级独立判断复合放大了误判率。收敛为"上游单点决策、下游纯执行"后，决策一致性和命中率均得到改善。具体SKILL清单是可插拔的业务数据层，与双Agent协同模式本身无关。', False)])

    add_body(doc, [('（3）人工入口动态管控', True)])
    add_body(doc, [('微观：', True), ('按"知识库答案/等待提示→技能答案→文档兜底"分阶段渐进交付答案，是否展示人工支持按钮由知识库答案状态与候选技能是否存在的组合判断是否已到达终态，具体组合逻辑如下：', False)])
    add_table(doc,
              ['知识库检索结果', '候选技能是否存在', '是否终态', '后续处理'],
              [
                  ['有答案', '不存在', '是', '直接展示知识库答案，提供人工支持按钮'],
                  ['有答案', '存在', '否', '知识库答案作为过渡参考，触发技能执行Agent，暂不提供人工支持按钮'],
                  ['无答案', '存在', '否', '展示等待提示，触发技能执行Agent，暂不提供人工支持按钮'],
                  ['无答案', '不存在', '否', '展示等待提示，跳过技能执行，直接进入文档推荐兜底，暂不提供人工支持按钮'],
              ])
    add_note(doc, '说明：技能执行完成、或文档推荐兜底返回后（无论结果如何），均视为到达终态，统一提供人工支持按钮，避免用户看到质量一般的中间结果就提前转向人工、也避免用户在不知情的情况下无限等待。')
    add_body(doc, [('宏观：', True), ('用户若1小时内未向机器人提问即直接尝试呼叫人工，会被拦截并引导先向机器人提问；熔断规则参数可配置，核心策略可迁移至其他场景。', False)])
    add_body(doc, [('这套机制天然构成可统计的问题漏斗（总提问数/自主解决数/转人工数），具体应用详见三、实践成效。', False)])

    add_body(doc, [('（4）结构化输出协议', True)])
    add_body(doc, [('配套设计了结构化Agent输出协议与具备异常兜底能力的解析逻辑，兼容大模型输出的格式异常与"未检索到相关信息"等无效占位答案，保障大模型生成的不确定性与确定性工作流编排之间的稳定对接。', False)])

    add_h2(doc, '2.5 知识自增长闭环：缺口识别与人工补齐')
    add_body(doc, [('机器尚无法自主判断"一次对话是否值得沉淀"、"该补知识库还是补技能"，这部分工作仍需人工完成。系统设计了两条独立互补的路径：', False)])
    add_body(doc, [('路径A·机器识别缺口：', True), ('复用核心引擎已计算出的知识库命中状态与候选技能是否存在这两个信号，将存在能力缺口的问题按组合情况分级收纳进问题池，供支持团队定期review针对性补齐（暂不做去重聚合，聚合去重是下一阶段规划，详见四、总结）。', False)])
    add_body(doc, [('路径B·人工主动沉淀：', True), ('人工支持过程中，支持人员可通过快捷指令将有价值的对话沉淀为知识库QA，或记录产品需求与BUG反馈。', False)])
    add_body(doc, [('两条路径分工明确：机器负责"发现缺口"，人工负责"决定并补齐"。', False)])

    # ============ 三、实践成效 ============
    add_h1(doc, '三、实践成效')
    add_body(doc, [('该方案本体已基于编码通(新)这一低代码技术支持场景完成落地，以下为落地后的具体情况。', False)])

    add_h2(doc, '3.1 问题拦截成效：问题漏斗')
    add_body(doc, [('引导熔断机制天然构成一个可统计的问题漏斗——总提问数、机器人自主解决数、转人工数之间形成清晰的漏斗关系，用统一口径衡量机器人实际拦截效果：', False)])
    add_table(doc,
              ['指标', '数值'],
              [
                  ['总提问数', '【待补充】'],
                  ['机器人自主解决数（拦截数）', '【待补充】'],
                  ['转人工支持数', '【待补充】'],
                  ['拦截率', '【待补充】'],
              ])

    add_h2(doc, '3.2 响应效率与人力成本')
    add_table(doc,
              ['指标', '上线前', '上线后'],
              [
                  ['平均响应/关单时长', '约50分钟', '【待补充】'],
                  ['日均人工支持投入', '4人力全天候', '【待补充】'],
                  ['日均问题处理量', '【待补充】', '【待补充】'],
              ])

    add_h2(doc, '3.3 知识沉淀能力建设情况')
    add_table(doc,
              ['能力项', '建设情况'],
              [
                  ['QA沉淀链路', '已打通，支持人员可通过快捷指令一键沉淀'],
                  ['需求/BUG反馈记录链路', '已打通，与QA沉淀共用同一快捷指令入口'],
                  ['Skill能力库', '已构建，覆盖7大业务模块，约30个SKILL'],
                  ['未解决问题池（分级收纳）', '已构建并自动运行，暂不做聚合去重'],
                  ['知识库QA条目数（累计）', '沉淀量仍在积累中'],
              ])
    add_note(doc, '说明：当前知识沉淀依赖支持人员的主观判断，存在触发门槛，团队已识别该环节为重点优化方向（详见四、总结的展望部分）。')

    add_h2(doc, '3.4 跨平台复用情况')
    add_body(doc, [('该通用架构已在编码通(新)之外的编码通2.0、飞码、MP等同一低代码产品矩阵内的平台复用，各平台只需替换知识库内容与SKILL清单即可直接复用核心架构，无需重新设计。当前复用范围仍局限于低代码产品矩阵内部，向其他业务线的推广是下一阶段的规划方向（详见四、总结的展望部分）。', False)])

    # ============ 四、总结 ============
    add_h1(doc, '四、总结')

    add_h2(doc, '4.1 创新性')
    add_bullet(doc, [('基于问题分层的三级递进应答引擎：', True), ('识别出确定性、推理性、场景可行性评估、需求反馈四类问题，设计知识库检索→Agent技能推理→文档兜底的逐层应答架构，精准匹配不同性质的问题，突破传统FAQ机器人"查得到才能答"的局限。', False)])
    add_bullet(doc, [('资源效率与用户体验的动态平衡机制：', True), ('通过知识库与候选技能状态组合驱动的终态判断，决定渐进交付每一步是否开放人工支持入口；叠加熔断时间窗口与提问有效性过滤两层宏观管控，在"计算资源利用效率"与"用户体验"之间建立可落地、可抗对抗性干扰的动态平衡。', False)])
    add_bullet(doc, [('双Agent单点决策架构：', True), ('技能推理环节由早期"上下游各自判断"收敛为"上游单点决策、下游纯执行"，这一调整由真实运行反馈驱动——识别到多级独立语义判断会复合放大误判率，收敛决策点后命中率与一致性均有改善。', False)])
    add_bullet(doc, [('知识自增长双轨闭环机制：', True), ('机器复用链路中已有信号自动识别知识库/技能覆盖缺口并分级收纳，人工在支持过程中主动快捷沉淀，两条路径互补，使知识库和技能库随系统运行持续自然生长。', False)])
    add_bullet(doc, [('通用引擎与可插拔数据层解耦的可复制架构：', True), ('核心引擎、Agent协同模式与业务数据层彻底解耦，已在编码通(新)之外的低代码产品矩阵内验证复用，具备规模化推广的现实基础。', False)])

    add_h2(doc, '4.2 效果和沉淀')
    add_body(doc, [('本方案已在"编码通(新)"平台正式上线运营，初步实现了对人工支持压力的有效分流：上线前平均关单时长约50分钟、日均需投入4人力全天候支持，上线后【待补充具体提升数据】。核心能力沉淀上，已构建覆盖7大业务模块、约30个SKILL的Agent推理能力库，打通了从人工支持到知识库的一键沉淀链路，以及机器自动识别缺口的问题池分级机制；该通用架构已进一步复用至编码通2.0、飞码、MP等同一低代码产品矩阵内的平台，初步验证了方案的跨平台可复制性。由于系统上线时间较短，知识沉淀的实际规模仍处于积累初期。', False)])

    add_h2(doc, '4.3 展望')
    add_bullet(doc, [('人工依赖范围的持续收窄：', True), ('当前"场景可行性评估"等复杂判断、以及知识/技能缺口的人工补齐，仍依赖人工完成——这是现阶段AI能力限制下的务实安排，而非设计终点。后续将持续跟踪大模型推理与判断能力的提升，逐步扩大机器自主覆盖的问题范围，收窄对人工的依赖。', False)])
    add_bullet(doc, [('问题分层数据的持续量化：', True), ('计划通过AI辅助分类对历史及新增咨询数据持续统计，使"问题分层"始终有数据驱动，反哺架构持续优化。', False)])
    add_bullet(doc, [('知识沉淀驱动机制优化：', True), ('当前人工主动沉淀依赖支持人员的主观判断，存在触发门槛。后续计划探索AI自动识别高价值对话并主动推荐沉淀，建立沉淀贡献的运营激励机制。', False)])
    add_bullet(doc, [('满意度数据埋点完善：', True), ('目前人工支持触发即视为机器人拦截失败，尚未对答案满意度做精细化埋点，后续计划补充显式反馈机制。', False)])
    add_bullet(doc, [('熔断策略丰富化：', True), ('当前熔断策略仅有"1小时内未提问不可直接呼叫人工"一条规则，后续计划结合更多场景（如同一问题重复追问次数）丰富熔断与引导策略。', False)])
    add_bullet(doc, [('人工交接体验优化：', True), ('当前转人工时机器人侧对话上下文不会自动同步给人工客服，后续计划打通上下文传递，避免用户重复描述问题。', False)])
    add_bullet(doc, [('问题池聚合去重与智能摘要推送：', True), ('当前问题池仅做分级、不做相似问题去重聚合。后续计划引入向量聚合、按频次/等级排序定期推送摘要，并探索缺口修复后的自动验证机制。', False)])
    add_bullet(doc, [('跨业务线推广：', True), ('当前复用范围仍局限于低代码产品矩阵内部，后续计划评估向其他业务线技术支持场景的适配可行性。', False)])

    out_path = '编码通新智能问答支持系统-创新奖申报材料.docx'
    doc.save(out_path)
    print('saved:', out_path)


if __name__ == '__main__':
    build()
